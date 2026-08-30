"""文件文本抽取：根据扩展名解析为纯文本。"""
import os


def extract_text(file_path: str) -> str:
    ext = os.path.splitext(file_path)[1].lower()
    if ext in (".txt", ".md"):
        with open(file_path, encoding="utf-8", errors="replace") as f:
            return f.read()
    if ext == ".pdf":
        return _extract_pdf(file_path)
    if ext == ".docx":
        return _extract_docx(file_path)
    if ext == ".pptx":
        return _extract_pptx(file_path)
    raise ValueError(f"不支持的文件类型: {ext}")


def _extract_pdf(file_path: str) -> str:
    from pypdf import PdfReader

    reader = PdfReader(file_path)
    return "\n".join((page.extract_text() or "") for page in reader.pages)


def _extract_docx(file_path: str) -> str:
    import docx

    doc = docx.Document(file_path)
    return "\n".join(p.text for p in doc.paragraphs)


def _extract_pptx(file_path: str) -> str:
    """逐页提取 PPTX 中文本框文字（含表格），按幻灯片组织。"""
    from pptx import Presentation

    prs = Presentation(file_path)
    parts = []
    for i, slide in enumerate(prs.slides, 1):
        lines: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs).strip()
                    if text:
                        lines.append(text)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        lines.append(" | ".join(cells))
        if lines:
            parts.append(f"## 幻灯片 {i}\n" + "\n".join(lines))
    return "\n\n".join(parts)
